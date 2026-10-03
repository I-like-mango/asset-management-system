import io
from decimal import Decimal, InvalidOperation

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from assets.models import Asset, AssetCategory
from assets.services.asset_service import create_asset, log_status_change
from common.constants import AssetStatus
from common.exceptions import BusinessError


IMPORT_HEADERS = [
    ("asset_no", "资产编号*"),
    ("name", "资产名称*"),
    ("category_code", "分类编码*"),
    ("brand", "品牌"),
    ("model", "型号"),
    ("serial_no", "序列号"),
    ("status", "状态(in_stock/loaned/maintenance/scrapped)"),
    ("location", "存放位置"),
    ("price", "采购价格"),
    ("remark", "备注"),
]

EXPORT_FIELD_MAP = {
    "asset_no": "资产编号",
    "name": "资产名称",
    "category": "分类",
    "brand": "品牌",
    "model": "型号",
    "serial_no": "序列号",
    "status": "状态",
    "location": "存放位置",
    "current_holder": "领用人",
    "loaned_at": "领用时间",
    "due_at": "到期时间",
    "price": "采购价格",
    "remark": "备注",
}


def build_import_template() -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产导入模板"
    ws.append([label for _, label in IMPORT_HEADERS])
    ws.append(["IT-001", "示例笔记本", "laptop", "Lenovo", "T14", "SN001", "in_stock", "仓库A", "8000", "示例"])
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="asset_import_template.xlsx"'
    return response


def import_assets_from_file(uploaded_file, operator):
    wb = load_workbook(uploaded_file, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise BusinessError("Excel 文件为空")
    header = [str(c).strip() if c else "" for c in rows[0]]
    expected = [label for _, label in IMPORT_HEADERS]
    if header[: len(expected)] != expected:
        raise BusinessError("模板表头不匹配，请下载最新模板")

    errors = []
    success = 0
    status_map = {v: k for k, v in AssetStatus.CHOICES}
    status_codes = {c[0] for c in AssetStatus.CHOICES}

    for idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        data = {IMPORT_HEADERS[i][0]: (row[i] if i < len(row) else "") for i in range(len(IMPORT_HEADERS))}
        asset_no = str(data["asset_no"] or "").strip()
        name = str(data["name"] or "").strip()
        cat_code = str(data["category_code"] or "").strip()
        if not asset_no or not name or not cat_code:
            errors.append(f"第{idx}行：资产编号、名称、分类编码为必填")
            continue
        if Asset.objects.filter(asset_no=asset_no, is_deleted=False).exists():
            errors.append(f"第{idx}行：资产编号 {asset_no} 已存在")
            continue
        category = AssetCategory.objects.filter(code=cat_code, is_active=True).first()
        if not category:
            errors.append(f"第{idx}行：分类编码 {cat_code} 不存在")
            continue
        status_raw = str(data.get("status") or AssetStatus.IN_STOCK).strip()
        if status_raw in status_map:
            status = status_raw
        elif status_raw in status_codes:
            status = status_raw
        else:
            errors.append(f"第{idx}行：无效状态 {status_raw}")
            continue
        price = None
        if data.get("price"):
            try:
                price = Decimal(str(data["price"]))
            except (InvalidOperation, ValueError):
                errors.append(f"第{idx}行：价格格式错误")
                continue
        try:
            create_asset(
                {
                    "asset_no": asset_no,
                    "name": name,
                    "category": category,
                    "brand": str(data.get("brand") or ""),
                    "model": str(data.get("model") or ""),
                    "serial_no": str(data.get("serial_no") or ""),
                    "status": status,
                    "location": str(data.get("location") or ""),
                    "price": price,
                    "remark": str(data.get("remark") or ""),
                },
                operator,
            )
            success += 1
        except BusinessError as e:
            errors.append(f"第{idx}行：{e.message}")

    return {"success": success, "errors": errors}


def export_assets_queryset(queryset, fields):
    if not fields:
        fields = list(EXPORT_FIELD_MAP.keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "资产导出"
    headers = [EXPORT_FIELD_MAP.get(f, f) for f in fields]
    ws.append(headers)
    status_labels = dict(AssetStatus.CHOICES)
    for asset in queryset.select_related("category", "current_holder"):
        row = []
        for f in fields:
            if f == "category":
                row.append(asset.category.name)
            elif f == "status":
                row.append(status_labels.get(asset.status, asset.status))
            elif f == "current_holder":
                row.append(asset.current_holder.username if asset.current_holder else "")
            elif f == "loaned_at":
                row.append(asset.loaned_at.strftime("%Y-%m-%d %H:%M") if asset.loaned_at else "")
            elif f == "due_at":
                row.append(asset.due_at.strftime("%Y-%m-%d %H:%M") if asset.due_at else "")
            else:
                row.append(getattr(asset, f, "") or "")
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="assets_export.xlsx"'
    return response
